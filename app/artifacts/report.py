from __future__ import annotations

import io
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from pptx import Presentation
from pptx.util import Inches, Pt


def _money(v: Any) -> str:
    return f"₹{float(v):,.2f}"


def _chart_buffer(fn, size=(6, 4)) -> io.BytesIO:
    buf = io.BytesIO()
    fn()
    plt.tight_layout()
    plt.savefig(buf, format="png", dpi=110)
    plt.close("all")
    buf.seek(0)
    return buf


def generate_analysis_pptx(snapshot: dict[str, Any], insights: str) -> bytes:
    """Build a weekly sales analytics deck. Returns PPTX bytes."""
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    summary = snapshot.get("sales_summary", {})
    daily = snapshot.get("daily_sales", [])
    top = snapshot.get("top_products", [])
    pay = snapshot.get("payment_breakdown", [])
    low = snapshot.get("low_stock", [])
    health = snapshot.get("stock_health", {})

    # ---- Title slide ----
    slide = prs.slides.add_slide(blank)
    tb = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1))
    tf = tb.text_frame
    tf.text = "Super Market — Sales Analytics"
    tf.paragraphs[0].font.size = Pt(40)
    sub = slide.shapes.add_textbox(Inches(1), Inches(3.6), Inches(11), Inches(1))
    stf = sub.text_frame
    stf.text = "Auto-generated report"
    stf.paragraphs[0].font.size = Pt(20)

    # ---- Summary slide ----
    slide = prs.slides.add_slide(blank)
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
    tf = tb.text_frame
    tf.text = "Sales Summary"
    tf.paragraphs[0].font.size = Pt(32)
    rows = [
        f"Total sales: {_money(summary.get('total_sales', 0))}",
        f"GST collected: {_money(summary.get('gst_collected', 0))}",
        f"Bills: {summary.get('bill_count', 0)}",
        f"SKUs: {health.get('total_skus', 0)}  |  Low stock: {health.get('low_stock_skus', 0)}  |  Out: {health.get('out_of_stock_skus', 0)}",
    ]
    tb2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(4))
    tf2 = tb2.text_frame
    tf2.text = "\n".join(rows)
    for p in tf2.paragraphs:
        p.font.size = Pt(20)

    # ---- Daily sales chart ----
    if daily:
        dates = [str(d["day"])[:10] for d in daily]
        totals = [float(d["total"]) for d in daily]
        buf = _chart_buffer(
            lambda: plt.plot(dates, totals, marker="o", color="#2f6fd6")
        )
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(buf, Inches(1), Inches(1.5), width=Inches(8))
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(10), Inches(1))
        tf = tb.text_frame
        tf.text = "Daily Sales"
        tf.paragraphs[0].font.size = Pt(32)

    # ---- Top products slide ----
    if top:
        names = [t["product_name"][:22] for t in top[:10]]
        revenue = [float(t["revenue"]) for t in top[:10]]
        buf = _chart_buffer(
            lambda: plt.barh(names[::-1], revenue[::-1], color="#4caf50")
        )
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(buf, Inches(1), Inches(1.5), width=Inches(8))
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(10), Inches(1))
        tf = tb.text_frame
        tf.text = "Top Products by Revenue"
        tf.paragraphs[0].font.size = Pt(32)

    # ---- Payment mix slide ----
    if pay:
        labels = [p["method"] for p in pay]
        values = [float(p["total"]) for p in pay]
        buf = _chart_buffer(lambda: plt.pie(values, labels=labels, autopct="%1.0f%%"))
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(buf, Inches(1), Inches(1.5), width=Inches(8))
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(10), Inches(1))
        tf = tb.text_frame
        tf.text = "Payment Mix"
        tf.paragraphs[0].font.size = Pt(32)

    # ---- Low stock slide ----
    slide = prs.slides.add_slide(blank)
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
    tf = tb.text_frame
    tf.text = "Stock Health"
    tf.paragraphs[0].font.size = Pt(32)
    rows = [f"{l['name']} — {float(l['quantity']):g} left (reorder {float(l['reorder_level']):g})" for l in low[:12]]
    tb2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(5))
    tf2 = tb2.text_frame
    tf2.text = "\n".join(rows) if rows else "All stock healthy."
    for p in tf2.paragraphs:
        p.font.size = Pt(18)

    # ---- Insights slide ----
    slide = prs.slides.add_slide(blank)
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
    tf = tb.text_frame
    tf.text = "AI Insights"
    tf.paragraphs[0].font.size = Pt(32)
    tb2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(5))
    tf2 = tb2.text_frame
    tf2.text = insights or "No insights available."
    tf2.word_wrap = True
    for p in tf2.paragraphs:
        p.font.size = Pt(16)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()